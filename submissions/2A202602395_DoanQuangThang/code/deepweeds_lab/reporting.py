"""Generate submission artifacts exclusively from persisted real experiment outputs."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from PIL import Image
from .dataset import CLASS_NAMES
from .train import write_json


def flat(row):
    return {k: v for k, v in row.items() if not isinstance(v, (dict, list))}


def mean_std(values):
    a = np.asarray(values, dtype=float)
    return float(a.mean()), float(a.std(ddof=1))


def md_table(df):
    return df.to_markdown(index=False, floatfmt='.4f')


def create_artifacts(root, backbones, training, inference, finals, selection):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    b = pd.DataFrame([flat(r) for r in backbones])
    t = pd.DataFrame([flat(r) for r in training])
    t['delta_macro_f1'] = t.macro_f1_val - t.loc[t.exp_id == 'T00', 'macro_f1_val'].iloc[0]
    i = pd.DataFrame([flat(r) for r in inference])
    f = pd.DataFrame([flat(r) for r in finals])
    aggregate, per_class, latency = [], [], []
    for tag in f.exp_id.unique():
        rows = [r for r in finals if r['exp_id'] == tag]
        summary = {'exp_id': tag, 'seed': 'mean ± sample std', 'n_seeds': len(rows)}
        for metric in ('macro_f1_val', 'macro_f1_test', 'top1_test', 'balanced_acc_test', 'ece_test'):
            mean, std = mean_std([r[metric] for r in rows])
            summary[metric] = mean; summary[metric + '_std'] = std
        aggregate.append(summary)
        for label, name in enumerate(CLASS_NAMES):
            pc = {'configuration': tag, 'class': name, 'label': label,
                  'support': rows[0]['per_class']['support'][label]}
            for metric in ('precision', 'recall', 'f1'):
                mean, std = mean_std([r['per_class'][metric][label] for r in rows])
                pc[metric] = mean; pc[metric + '_std'] = std
            per_class.append(pc)
    for r in finals:
        latency.append(dict(configuration=r['exp_id'], seed=r['seed'], **r['latency']))
    for r in inference:
        latency.append(dict(configuration=r['exp_id'], **{k: r[k] for k in
            ('gpu', 'dtype', 'batch', 'img_size', 'p50', 'p95', 'p99', 'images_per_s', 'torch', 'preprocessing', 'fused_bn')}))
        latency.append(dict(configuration=r['exp_id'], **r['latency_batch32']))
    summary_rows = []
    for r in backbones + training:
        summary_rows.append(dict(exp_id=r['exp_id'], stage='training', backbone=r['backbone'],
                                 macro_f1_val=r['macro_f1_val'], top1_val=r['top1_val'],
                                 latency_ms=r.get('latency_ms', np.nan)))
    for r in inference:
        summary_rows.append(dict(exp_id=r['exp_id'], stage='inference', backbone=selection['backbone'],
                                 macro_f1_val=r['macro_f1'], top1_val=r['top1'], latency_ms=r['p95']))
    summary_frame = pd.DataFrame(summary_rows).sort_values('macro_f1_val', ascending=False).head(10)
    final_frame = pd.concat([f, pd.DataFrame(aggregate)], ignore_index=True)
    with pd.ExcelWriter(root / 'results.xlsx', engine='openpyxl') as writer:
        sheets = {'Backbones': b, 'Training': t, 'Inference': i, 'Final': final_frame,
                  'PerClass': pd.DataFrame(per_class), 'Latency': pd.DataFrame(latency), 'Summary': summary_frame}
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]; sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in sheet[1]:
                cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='DCEAF7')
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(44, max(14, len(str(column[0].value)) + 2))
                for cell in column[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = '0.0000'
            metric_col = 'macro_f1_val' if 'macro_f1_val' in frame else 'macro_f1' if 'macro_f1' in frame else None
            if metric_col and len(frame):
                index = int(frame[metric_col].astype(float).idxmax()) + 2
                for cell in sheet[index]:
                    cell.fill = PatternFill('solid', fgColor='E2F0D9')
    fig, ax = plt.subplots(figsize=(8, 5))
    for r in inference:
        ax.scatter(r['p95'], r['macro_f1']); ax.annotate(r['exp_id'] + ' ' + r['method'], (r['p95'], r['macro_f1']), fontsize=8)
    ax.set_xlabel('Batch-1 p95 latency (ms), preprocessing excluded'); ax.set_ylabel('Validation macro-F1'); ax.grid(alpha=.2)
    fig.tight_layout(); fig.savefig(root / 'accuracy_latency.png', dpi=150); plt.close(fig)
    final_rows = [r for r in finals if r['exp_id'] == 'F01']
    cm = np.asarray([r['confusion'] for r in final_rows]).mean(0)
    fig, ax = plt.subplots(figsize=(9, 8)); im = ax.imshow(cm, cmap='Blues'); fig.colorbar(im, ax=ax)
    ax.set_xticks(range(9), CLASS_NAMES, rotation=60, ha='right'); ax.set_yticks(range(9), CLASS_NAMES)
    ax.set_xlabel('Predicted'); ax.set_ylabel('True'); ax.set_title('Final test confusion: mean counts across seeds')
    for y in range(9):
        for x in range(9):
            ax.text(x, y, f'{cm[y,x]:.1f}', ha='center', va='center', fontsize=7,
                    color='white' if cm[y,x] > cm.max() / 2 else 'black')
    fig.tight_layout(); fig.savefig(root / 'confusion_matrix.png', dpi=150); plt.close(fig)
    seed = selection['seeds'][0]
    predictions = pd.read_csv(root / 'predictions' / f'F01_seed{seed}_test.csv')
    wrong = predictions[predictions.y_true != predictions.y_pred]
    hard = wrong[wrong.y_true.isin([0, 7]) & wrong.y_pred.isin([0, 7])]
    examples = pd.concat([hard, wrong]).drop_duplicates('Filename').head(12)
    write_json(root / 'error_examples.json', examples[['Filename', 'y_true', 'y_pred']].to_dict('records'))
    if len(examples):
        fig, axes = plt.subplots(3, 4, figsize=(12, 9))
        for ax in axes.flat:
            ax.axis('off')
        for ax, (_, row) in zip(axes.flat, examples.iterrows()):
            with Image.open(Path(selection['training_config']['images_dir']) / row.Filename) as image:
                ax.imshow(image.convert('RGB'))
            ax.set_title(f'{CLASS_NAMES[int(row.y_true)]}\n→ {CLASS_NAMES[int(row.y_pred)]}', fontsize=8)
        fig.tight_layout(); fig.savefig(root / 'misclassified_images.png', dpi=130); plt.close(fig)
    ag = {r['exp_id']: r for r in aggregate}; final, baseline = ag['F01'], ag['T00']
    delta = final['macro_f1_test'] - baseline['macro_f1_test']
    noise = max(final['macro_f1_test_std'], baseline['macro_f1_test_std'])
    conclusion = 'Mức cải thiện lớn hơn độ lệch chuẩn lớn nhất của hai nhóm.' if delta > noise else 'Chênh lệch chưa vượt nhiễu giữa các seed; chưa đủ bằng chứng khẳng định tốt hơn.'
    effect_backbone = b.macro_f1_val.max() - b.macro_f1_val.min()
    effect_training = t.macro_f1_val.max() - t.loc[t.exp_id == 'T00', 'macro_f1_val'].iloc[0]
    effect_inference = i.macro_f1.max() - i.loc[i.method == 'single', 'macro_f1'].iloc[0]
    effects = {'chọn backbone (khoảng tốt nhất–kém nhất)': effect_backbone,
               'công thức huấn luyện (so với nền)': effect_training,
               'suy luận (so với một ảnh)': effect_inference}
    winner_effect = max(effects, key=effects.get)
    gaps = [abs(r['macro_f1_val'] - r['macro_f1_test']) for r in final_rows]
    cal_before = np.mean([r['ece_uncal_test'] for r in final_rows]); cal_after = final['ece_test']
    off = cm.copy(); np.fill_diagonal(off, 0); true, pred = np.unravel_index(off.argmax(), off.shape)
    infer_choice = next(r for r in inference if r['method'] == selection['inference'])
    rt = ag.get('R01')
    realtime = (f"Cấu hình R01 dùng {selection.get('realtime_method')}; p95 trên val = "
                f"{np.mean([r['p95_ms'] for r in finals if r['exp_id']=='R01']):.2f} ms (trung bình đo qua seed), "
                f"macro-F1 test = {rt['macro_f1_test']:.4f} ± {rt['macro_f1_test_std']:.4f}.") if rt else 'Chưa có cấu hình đo được p95 ≤100 ms; không tuyên bố đáp ứng thời gian thực.'
    if rt and np.mean([r['p95_ms'] for r in finals if r['exp_id']=='R01']) > 100:
        realtime += ' Đo vòng cuối vượt 100 ms; chưa đạt ngân sách thời gian thực.'
    text = f'''# Báo cáo Lab Day 2 — DeepWeeds

## Tóm tắt
Đã so sánh {len(backbones)} backbone, {len(training)} công thức huấn luyện và {len(inference)-1} phương pháp suy luận ngoài mốc.
Cấu hình F01: **{selection['backbone']}**, công thức **{selection['training_config']['exp_id']}**, suy luận **{selection['inference']}**, kèm temperature scaling khớp riêng trên val cho mỗi seed.
Macro-F1 test: **{final['macro_f1_test']:.4f} ± {final['macro_f1_test_std']:.4f}**; top-1 test: **{final['top1_test']:.4f} ± {final['top1_test_std']:.4f}**.
Mức cải thiện macro-F1 so với T00: {delta:+.4f}. {conclusion}

## Dữ liệu và thiết lập
Dùng nguyên CSV fold 0 của tác giả; train/val/test không giao nhau và đủ 17.509 ảnh. Số đếm thật nằm trong `split_checks.json`. Chỉ dùng train cho trọng số; val cho mọi lựa chọn và nhiệt độ; test chỉ ở vòng cuối sau `selection_locked.json`.
![Phân bố lớp](class_distribution.png)
![Ảnh mẫu train](sample_images.png)
Kiểm tra pipeline với mạng tuyến tính nhỏ, độc lập với kết quả backbone, nằm trong `pipeline_checks.json`; ảnh sau augmentation ở `augmented_images.png`.
Nền: ImageNet finetune, ảnh 224, crop/lật ngang, AdamW, LR backbone 1e-4/head 1e-3, decay 0.05 (trừ norm/bias), warmup một epoch rồi cosine theo bước, CE, {selection['training_config']['epochs']} epoch, batch {selection['training_config']['batch_size']}, AMP. Tiền xử lý dùng mean/std đúng trọng số; val/test resize 256 rồi crop 224. GMAC từ fvcore, phép tính không được hỗ trợ được ghi trong log.
GPU: {infer_choice['gpu']}; torch: {infer_choice['torch']}. Phiên bản đầy đủ trong `environment.json`. Seed vòng cuối: {selection['seeds']}; std mẫu ddof=1. Vòng sàng và ablation chỉ một seed.

## So sánh backbone
{md_table(b[['exp_id','backbone','pretrained_tag','params_m','gmac','macro_f1_val','top1_val','train_seconds_per_epoch','latency_ms']])}
Chọn {selection['backbone']} vì macro-F1 val cao nhất trong vòng sàng. Bảng vẫn ghi tốc độ và kích thước để thể hiện chi phí; lựa chọn này ưu tiên chất lượng.

## Công thức huấn luyện
{md_table(t[['exp_id','axis','changes','macro_f1_val','top1_val','delta_macro_f1']])}
Mỗi T01–T08 chỉ đổi một trục so với T00. T09 là kết hợp thử nghiệm; lựa chọn cuối dựa trên macro-F1 val, không mặc định kết hợp sẽ tốt hơn.
Các delta một seed là dấu hiệu sàng lọc, chưa phải bằng chứng thống kê. Vòng cuối mới so sánh với nhiễu đa seed.
EMA trung bình tham số và sao chép buffer BN từ mạng đang huấn luyện; checkpoint được đánh giá bằng bản EMA. Seed đầu của vòng cuối tái sử dụng checkpoint của đúng công thức đã chạy, nguồn được ghi trong `Final`.

## Suy luận, hiệu chuẩn và độ trễ
{md_table(i[['exp_id','method','K','macro_f1','top1','ece','p50','p95','p99','relative_cost']])}
![Đánh đổi độ chính xác–độ trễ](accuracy_latency.png)
Độ trễ đo với 10 lượt warmup, 50 lượt có đồng bộ GPU, batch 1 và 32, gồm forward/softmax/gộp view, không gồm đọc ảnh và tiền xử lý. TTA và AMP được đo thực tế.
ECE test trước hiệu chuẩn: {cal_before:.4f}; sau: {cal_after:.4f}. Tối ưu temperature theo NLL trên val không đảm bảo giảm ECE trên test; số liệu trên là kết quả thực tế.

## Chung kết và phân tích lỗi
{md_table(pd.DataFrame(aggregate)[['exp_id','n_seeds','macro_f1_val','macro_f1_test','macro_f1_test_std','top1_test','top1_test_std','ece_test']])}
![Ma trận nhầm lẫn](confusion_matrix.png)
{md_table(pd.DataFrame(per_class)[['configuration','class','support','precision','recall','f1']])}
Cặp nhầm có số ảnh trung bình lớn nhất: {CLASS_NAMES[true]} → {CLASS_NAMES[pred]} ({off[true,pred]:.1f} ảnh). Cần xem ảnh lỗi trước khi kết luận nguyên nhân; tương đồng hình dạng và bối cảnh là giả thuyết, chưa được kiểm chứng nhân quả.
![Ảnh dự đoán sai](misclassified_images.png)
Các ví dụ ưu tiên cặp Chinee Apple ↔ Snake Weed. Tên ảnh và nhãn nằm trong `error_examples.json`.
Chênh lệch macro-F1 val/test tuyệt đối từng seed: {gaps}. Có thể khác biệt do độ khó mẫu và chọn checkpoint trên val; không điều chỉnh lại cấu hình sau khi đọc test.

## Kết luận và khuyến nghị
{conclusion}
Khoảng chênh trên val: backbone {effect_backbone:.4f}; huấn luyện {effect_training:.4f}; suy luận {effect_inference:.4f}. Trong các phép so sánh đã chạy, **{winner_effect}** có khoảng lớn nhất. Các khoảng này có mốc khác nhau và chỉ một seed, nên không coi là phân rã nhân quả hay bằng chứng phổ quát.
{realtime}
TTA nhiều view tốn thêm forward; dùng bảng độ trễ thực tế để quyết định chạy ngoại tuyến hay trên robot.

## Hạn chế và việc tiếp theo
Chỉ dùng một fold, các vòng sàng một seed và ngân sách {selection['training_config']['epochs']} epoch. Split ngẫu nhiên không theo địa điểm có thể cho test lạc quan khi triển khai nơi mới. Seed cố định và cuDNN deterministic không bảo đảm giống tuyệt đối giữa phần cứng; khi resume, trạng thái worker augmentation có thể khác phiên liên tục.
Các thất bại nếu có được lưu trong `inference_failures.json`. Không dùng số tham khảo bài báo làm kết quả của mình. Bài báo gốc huấn luyện lâu hơn; chưa kiểm chứng các miền mùa/ánh sáng khác, ONNX hay robot thật.

## Phụ lục
Cấu hình mọi lần chạy: `runs/<exp_id>/seed<k>/config.json`; log: `history.csv`; trọng số tốt nhất: `best.pt`; biểu đồ: `curves/`; dự đoán: `predictions/`. Nguồn trọng số và tag: `pretrained.json`. Đánh giá đối chiếu: `eval_score_*.txt`, `eval_grade.txt`, `eval_out/`.
Notebook: https://www.kaggle.com/code/thngonquang/deepweeds-day2
'''
    (root / 'report.md').write_text(text, encoding='utf-8')
    (root / 'README.md').write_text('''# DeepWeeds Day2 — kết quả chạy thật
Notebook: https://www.kaggle.com/code/thngonquang/deepweeds-day2

Đọc report.md và results.xlsx; cấu hình đã chốt nằm trong selection_locked.json. Lệnh chạy lại: `python -m deepweeds_lab.experiments --output lab_output --epochs 10 --batch-size 32`, từ môi trường có package và eval.py trong PYTHONPATH. Phiên bản ở environment.json; seed 0,1,2. Giữ nguyên config khi resume, không đổi split hoặc dùng test để chọn lại model. Checkpoint và dataset chỉ lưu ngoài Git.
''', encoding='utf-8')
